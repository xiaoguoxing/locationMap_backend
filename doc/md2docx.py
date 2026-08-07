# -*- coding: utf-8 -*-
"""
按 tools:word-format skill 规范，将 Markdown 转为 Word docx。
中文黑体/宋体/楷体，西文 Times New Roman，代码块 Consolas，表格/引用/图片完整支持。
自动生成封面页（标题 + 文档信息表格）+ 目录页（真实 Word TOC 域，可右键更新）。
"""
import os
import re
import datetime
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

FONT_LATIN = 'Times New Roman'
FONT_BODY_CN = '宋体'
FONT_HEADING_CN = '黑体'
FONT_QUOTE_CN = '楷体'
FONT_CODE = 'Consolas'

# 封面页表格的字段顺序（项目列固定文案；内容列从 metadata 取）
COVER_FIELDS = [
    ('文档名称', 'title'),
    ('版本',     'version'),
    ('编制单位', 'organization'),
    ('编制日期', 'date'),
]

# 默认元数据（找不到 frontmatter 且启发式提取失败时的兜底）
COVER_DEFAULTS = {
    'version': 'V1.0',
}

# 识别 md 中"已存在的目录段"的标题（中英文都覆盖，跳过避免与 Word 真实 TOC 重复）
TOC_HEADING_PATTERN = re.compile(r'^\s*(目\s*录|Contents|Table\s+of\s+Contents|TOC)\s*$', re.IGNORECASE)


def _clear_theme_fonts(rFonts):
    for attr in ('asciiTheme', 'hAnsiTheme', 'eastAsiaTheme', 'cstheme'):
        key = qn(f'w:{attr}')
        if rFonts.get(key) is not None:
            del rFonts.attrib[key]


def set_run_fonts(run, east_asia=FONT_BODY_CN, latin=FONT_LATIN):
    run.font.name = latin
    rPr = run.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), latin)
    rFonts.set(qn('w:hAnsi'), latin)
    rFonts.set(qn('w:eastAsia'), east_asia)
    _clear_theme_fonts(rFonts)


def set_run_code_font(run):
    run.font.name = FONT_CODE
    rPr = run.element.get_or_add_rPr()
    rFonts = rPr.find(qn('w:rFonts'))
    if rFonts is None:
        rFonts = OxmlElement('w:rFonts')
        rPr.insert(0, rFonts)
    rFonts.set(qn('w:ascii'), FONT_CODE)
    rFonts.set(qn('w:hAnsi'), FONT_CODE)
    rFonts.set(qn('w:eastAsia'), FONT_CODE)
    _clear_theme_fonts(rFonts)


def configure_styles(doc):
    section = doc.sections[0]
    section.page_width = Cm(21)
    section.page_height = Cm(29.7)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(3.17)
    section.right_margin = Cm(3.17)

    normal = doc.styles['Normal']
    normal.font.name = FONT_LATIN
    normal.font.size = Pt(12)
    normal.paragraph_format.line_spacing = 1.5
    normal.paragraph_format.first_line_indent = Cm(0.74)
    normal.paragraph_format.space_before = Pt(3)
    normal.paragraph_format.space_after = Pt(3)
    normal.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_BODY_CN)

    heading_config = [
        (1, Pt(22), WD_ALIGN_PARAGRAPH.CENTER, Pt(24), Pt(12), RGBColor(0, 0, 0)),
        (2, Pt(16), WD_ALIGN_PARAGRAPH.LEFT, Pt(18), Pt(8), RGBColor(0, 0, 0)),
        (3, Pt(14), WD_ALIGN_PARAGRAPH.LEFT, Pt(12), Pt(6), RGBColor(0, 0, 0)),
        (4, Pt(12), WD_ALIGN_PARAGRAPH.LEFT, Pt(8), Pt(4), RGBColor(0x33, 0x33, 0x33)),
    ]
    for level, size, align, before, after, color in heading_config:
        h = doc.styles[f'Heading {level}']
        h.font.name = FONT_LATIN
        h.font.size = size
        h.font.bold = True
        h.font.color.rgb = color
        h.paragraph_format.alignment = align
        h.paragraph_format.space_before = before
        h.paragraph_format.space_after = after
        h.paragraph_format.first_line_indent = Cm(0)
        h.paragraph_format.left_indent = Cm(0)
        h.element.rPr.rFonts.set(qn('w:eastAsia'), FONT_HEADING_CN)


def _emit_run(paragraph, text, *, bold, east_asia, size=None, code=False, color=None):
    run = paragraph.add_run(text)
    run.bold = bold
    if size:
        run.font.size = size
    if color:
        run.font.color.rgb = color
    if code:
        set_run_code_font(run)
    else:
        set_run_fonts(run, east_asia=east_asia)
    return run


def add_inline_text(paragraph, text, *, east_asia=FONT_BODY_CN, size=None, bold=False, color=None):
    """处理 **bold** / `code` 行内标记"""
    for bold_part in re.split(r'(\*\*.*?\*\*)', text):
        if not bold_part:
            continue
        is_bold = bold
        content = bold_part
        if bold_part.startswith('**') and bold_part.endswith('**') and len(bold_part) > 4:
            is_bold = True
            content = bold_part[2:-2]
        for code_part in re.split(r'(`[^`]+`)', content):
            if not code_part:
                continue
            if code_part.startswith('`') and code_part.endswith('`') and len(code_part) > 2:
                _emit_run(paragraph, code_part[1:-1], bold=is_bold,
                          east_asia=east_asia, size=size, code=True, color=color)
            else:
                _emit_run(paragraph, code_part, bold=is_bold,
                          east_asia=east_asia, size=size, code=False, color=color)


def add_heading(doc, text, level):
    h = doc.add_heading(level=level)
    h.paragraph_format.first_line_indent = Cm(0)
    add_inline_text(h, text, east_asia=FONT_HEADING_CN, bold=True)
    for run in h.runs:
        set_run_fonts(run, east_asia=FONT_HEADING_CN)
    return h


def add_paragraph(doc, text, *, bullet=False, number=None, checkbox=None, indent_level=0):
    p = doc.add_paragraph()
    if bullet or number is not None or checkbox is not None:
        p.paragraph_format.first_line_indent = Cm(0)
        p.paragraph_format.left_indent = Cm(0.74 * max(1, indent_level))
        prefix = ''
        if checkbox is not None:
            prefix = '☑ ' if checkbox else '☐ '
        elif bullet:
            prefix = '• '
        elif number is not None:
            prefix = f'{number}. '
        if prefix:
            _emit_run(p, prefix, bold=False, east_asia=FONT_BODY_CN)
    add_inline_text(p, text, east_asia=FONT_BODY_CN)


def add_blockquote(doc, lines):
    text = ' '.join(lines)
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1.5)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.5
    pPr = p._p.get_or_add_pPr()
    pBdr = parse_xml(
        f'<w:pBdr {nsdecls("w")}>'
        f'<w:left w:val="single" w:sz="12" w:space="8" w:color="CCCCCC"/>'
        f'</w:pBdr>'
    )
    pPr.append(pBdr)
    add_inline_text(p, text, east_asia=FONT_QUOTE_CN,
                    size=Pt(11), color=RGBColor(0x66, 0x66, 0x66))


def add_code_block(doc, code_text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Cm(1)
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.space_before = Pt(6)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.line_spacing = 1.15
    pPr = p._p.get_or_add_pPr()
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="F5F5F5"/>')
    pPr.append(shading)
    run = p.add_run(code_text)
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(0x33, 0x33, 0x33)
    set_run_code_font(run)


def _setup_cell(cell, text, *, is_header=False):
    cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
    cell.text = ''
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    east_asia = FONT_HEADING_CN if is_header else FONT_BODY_CN
    add_inline_text(p, text.strip(), east_asia=east_asia, size=Pt(10), bold=is_header)
    if is_header:
        shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="E8E8E8"/>')
        cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc, headers, rows):
    if not headers:
        return
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = True

    for i, h in enumerate(headers):
        _setup_cell(table.rows[0].cells[i], h, is_header=True)
    for ri, row in enumerate(rows):
        for ci in range(len(headers)):
            val = row[ci] if ci < len(row) else ''
            _setup_cell(table.rows[ri + 1].cells[ci], val)
    doc.add_paragraph()


def add_image(doc, img_path, max_width_cm=15.5):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    if not os.path.exists(img_path):
        _emit_run(p, f'[图片丢失: {img_path}]', bold=False, east_asia=FONT_BODY_CN)
        return
    run = p.add_run()
    run.add_picture(img_path, width=Cm(max_width_cm))


# ---- 封面页 / 目录页 / 元数据 ----

def parse_frontmatter(lines):
    """
    解析 md 文件头部的 YAML frontmatter（如有）：
        ---
        title: 信创技术体系...
        version: V1.0
        organization: 深圳市科荣软件股份有限公司
        date: 2026年5月
        ---

    返回 (metadata_dict, body_lines)，body_lines 是去掉 frontmatter 后的剩余行。
    没找到 frontmatter 时返回 ({}, lines)。
    """
    if not lines or lines[0].rstrip('\n').strip() != '---':
        return {}, lines
    end = None
    for i in range(1, len(lines)):
        if lines[i].rstrip('\n').strip() == '---':
            end = i
            break
    if end is None:
        return {}, lines
    metadata = {}
    for raw in lines[1:end]:
        line = raw.rstrip('\n').strip()
        if not line or line.startswith('#'):
            continue
        m = re.match(r'^([\w-]+)\s*:\s*(.+)$', line)
        if m:
            key = m.group(1).strip().lower()
            val = m.group(2).strip().strip('"').strip("'")
            metadata[key] = val
    return metadata, lines[end + 1:]


def heuristic_metadata(lines):
    """
    没 frontmatter 时启发式提取元数据：
    - title: 第一个 H1（# xxx）
    - organization: 第一个粗体段（**xxx**）
    - date: 第二个粗体段
    扫描范围限制在前 30 行，避免误吞正文。
    """
    meta = {}
    bold_count = 0
    for raw in lines[:30]:
        line = raw.rstrip('\n').strip()
        if not line:
            continue
        if 'title' not in meta:
            m = re.match(r'^#\s+(.+)$', line)
            if m:
                meta['title'] = m.group(1).strip()
                continue
        m = re.match(r'^\*\*(.+?)\*\*$', line)
        if m:
            bold_count += 1
            if bold_count == 1:
                meta['organization'] = m.group(1).strip()
            elif bold_count == 2:
                meta['date'] = m.group(1).strip()
                break
    return meta


def resolve_metadata(lines):
    """三级回退：frontmatter > 启发式提取 > 默认值。返回 (metadata, body_lines)"""
    meta_fm, body = parse_frontmatter(lines)
    meta_heu = heuristic_metadata(body)
    final = dict(COVER_DEFAULTS)
    final.update(meta_heu)
    final.update(meta_fm)  # frontmatter 覆盖启发式
    if 'date' not in final:
        final['date'] = datetime.datetime.now().strftime('%Y年%-m月') if os.name != 'nt' \
                       else f"{datetime.datetime.now().year}年{datetime.datetime.now().month}月"
    if 'title' not in final:
        final['title'] = '未命名文档'
    if 'organization' not in final:
        final['organization'] = ''
    return final, body


def add_page_break(doc):
    """在文档末尾添加强制分页符"""
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.add_run().add_break(WD_BREAK.PAGE)


def add_cover_page(doc, metadata):
    """
    封面页：大标题居中 + 大段空白 + 文档信息表格（项目/内容两列，4 行）
    metadata 必须含 title/version/organization/date 四键（resolve_metadata 已兜底）
    ⚠️ 封面标题段**不**使用 Heading 1 样式，避免被 TOC 域识别成"目录的第一条"。
    """
    # 大标题：普通段落 + 28pt 加粗居中（不是 Heading 1，否则会进入 TOC）
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    p.paragraph_format.space_before = Pt(80)
    p.paragraph_format.space_after = Pt(40)
    p.paragraph_format.line_spacing = 1.5
    run = p.add_run(metadata['title'])
    run.bold = True
    run.font.size = Pt(28)
    set_run_fonts(run, east_asia=FONT_HEADING_CN)

    # 标题与表格之间的空段（半页空白）
    for _ in range(8):
        sp = doc.add_paragraph()
        sp.paragraph_format.first_line_indent = Cm(0)
        sp.paragraph_format.line_spacing = 1.5

    # 文档信息表格（2 列：项目 / 内容；4 行数据）
    headers = ['项目', '内容']
    rows = [[label, str(metadata.get(key, ''))] for label, key in COVER_FIELDS]
    add_table(doc, headers, rows)

    add_page_break(doc)


def add_toc_page(doc):
    """
    目录页：真实 Word TOC 域，用户首次打开 docx 时 Word/WPS 会提示"更新域"，
    或右键目录 → 更新域 → 更新整个目录，即可填充真实目录条目（基于文档中的 Heading 1-3）。
    TOC 开关说明：
      \\o "1-3"  显示一级到三级标题
      \\h        生成超链接（点击跳转到对应章节）
      \\z        在 Web 视图下隐藏制表符前导符
      \\u        使用大纲级别构建目录
    """
    # 目录标题
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.first_line_indent = Cm(0)
    p_title.paragraph_format.space_before = Pt(24)
    p_title.paragraph_format.space_after = Pt(18)
    run_title = p_title.add_run('目  录')
    run_title.bold = True
    run_title.font.size = Pt(22)
    set_run_fonts(run_title, east_asia=FONT_HEADING_CN)

    # TOC 域占位段
    p = doc.add_paragraph()
    p.paragraph_format.first_line_indent = Cm(0)
    p.paragraph_format.left_indent = Cm(0)
    run = p.add_run()

    fldChar_begin = OxmlElement('w:fldChar')
    fldChar_begin.set(qn('w:fldCharType'), 'begin')

    instrText = OxmlElement('w:instrText')
    instrText.set(qn('xml:space'), 'preserve')
    instrText.text = r'TOC \o "1-3" \h \z \u'

    fldChar_sep = OxmlElement('w:fldChar')
    fldChar_sep.set(qn('w:fldCharType'), 'separate')

    placeholder = OxmlElement('w:t')
    placeholder.text = '（首次打开请右键此处 → 更新域 → 更新整个目录）'

    fldChar_end = OxmlElement('w:fldChar')
    fldChar_end.set(qn('w:fldCharType'), 'end')

    r_element = run._r
    r_element.append(fldChar_begin)
    r_element.append(instrText)
    r_element.append(fldChar_sep)
    r_element.append(placeholder)
    r_element.append(fldChar_end)

    # 提示 Word 在打开时自动更新所有域（settings.xml 加 <w:updateFields w:val="true"/>）
    settings = doc.settings.element
    update_fields = settings.find(qn('w:updateFields'))
    if update_fields is None:
        update_fields = OxmlElement('w:updateFields')
        update_fields.set(qn('w:val'), 'true')
        settings.append(update_fields)

    add_page_break(doc)


def _parse_table_row(line):
    return [c.strip() for c in line.strip().strip('|').split('|')]


_TABLE_SEP_RE = re.compile(r'^\|\s*:?-+:?\s*(\|\s*:?-+:?\s*)+\|?\s*$')


def convert(md_path, docx_path, *, with_cover=True, with_toc=True, metadata_override=None):
    """
    主转换函数。
      with_cover       是否生成封面页（默认 True）
      with_toc         是否生成目录页（默认 True，使用 Word 真实 TOC 域）
      metadata_override 调用方可显式传入元数据 dict（key: title/version/organization/date），
                      会覆盖 frontmatter / 启发式提取的结果
    """
    doc = Document()
    configure_styles(doc)

    base_dir = os.path.dirname(os.path.abspath(md_path))
    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    # 解析元数据（frontmatter > 启发式 > 默认值），生成封面页 + 目录页
    metadata, body_lines = resolve_metadata(lines)
    if metadata_override:
        metadata.update(metadata_override)

    if with_cover:
        add_cover_page(doc, metadata)
    if with_toc:
        add_toc_page(doc)

    # 正文解析（跳过启发式提取已用过的标题/粗体行 + md 中已存在的"## 目录"段）
    lines = body_lines

    in_code = False
    code_buf = []
    table_headers = None
    table_rows = []
    quote_buf = []
    # 跳过 md 中已有"## 目录"段（避免与 Word 真实 TOC 重复）
    skip_toc_section = False
    skip_toc_level = 0
    # 跳过封面已用的元数据行（H1 标题 + 前两个粗体段 + 第一条 ---）
    cover_h1_skipped = (not with_cover)  # 不生成封面则不跳
    cover_bolds_to_skip = 0 if (not with_cover) else 2
    cover_separator_skipped = (not with_cover)
    cover_metadata_phase = with_cover  # True 期间扫描封面元数据行（H1 + 粗体 + ---）

    def flush_table():
        nonlocal table_headers, table_rows
        if table_headers is not None:
            add_table(doc, table_headers, table_rows)
            table_headers = None
            table_rows = []

    def flush_quote():
        nonlocal quote_buf
        if quote_buf:
            add_blockquote(doc, quote_buf)
            quote_buf = []

    for raw in lines:
        line = raw.rstrip('\n')
        stripped = line.strip()

        # ---- 跳过封面已用元数据（仅在文件开头若干行）----
        if cover_metadata_phase:
            if not stripped:
                continue
            if (not cover_h1_skipped) and re.match(r'^#\s+', stripped):
                cover_h1_skipped = True
                continue
            if cover_h1_skipped and cover_bolds_to_skip > 0 and re.match(r'^\*\*.+?\*\*$', stripped):
                cover_bolds_to_skip -= 1
                continue
            if cover_h1_skipped and cover_bolds_to_skip == 0 and (not cover_separator_skipped) and stripped == '---':
                cover_separator_skipped = True
                cover_metadata_phase = False  # 元数据段结束，进入正文
                continue
            # 一旦遇到非元数据格式的内容（如直接进入 ## 标题），结束元数据扫描
            if cover_h1_skipped and (cover_bolds_to_skip < 2 or re.match(r'^#{1,6}\s+', stripped)):
                cover_metadata_phase = False
                # 不 continue，让本行进入正常解析

        if stripped.startswith('```'):
            if in_code:
                add_code_block(doc, '\n'.join(code_buf))
                code_buf = []
                in_code = False
            else:
                flush_table()
                flush_quote()
                in_code = True
            continue
        if in_code:
            code_buf.append(line)
            continue

        if line.startswith('|') and '|' in line[1:]:
            if _TABLE_SEP_RE.match(stripped):
                continue
            parts = _parse_table_row(line)
            if table_headers is None:
                flush_quote()
                table_headers = parts
            else:
                table_rows.append(parts)
            continue
        flush_table()

        if stripped in ('---', '***', '___'):
            flush_quote()
            continue

        m = re.match(r'^(#{1,4})\s+(.+)$', stripped)
        if m:
            heading_text = m.group(2).strip()
            heading_level = len(m.group(1))
            # 进入"目录段"（## 目录 / Contents 等）→ 设标志，跳过该段
            if TOC_HEADING_PATTERN.match(heading_text):
                skip_toc_section = True
                skip_toc_level = heading_level
                continue
            # 退出"目录段"：遇到同级或更高级标题
            if skip_toc_section and heading_level <= skip_toc_level:
                skip_toc_section = False
            if skip_toc_section:
                continue
            flush_quote()
            add_heading(doc, heading_text, heading_level)
            continue

        # 在"目录段"内的非标题行（如手写目录条目 - [xxx](#yyy)）也跳过
        if skip_toc_section:
            continue

        m = re.match(r'^!\[([^\]]*)\]\(([^)]+)\)\s*$', stripped)
        if m:
            flush_quote()
            img_path = os.path.normpath(os.path.join(base_dir, m.group(2)))
            add_image(doc, img_path)
            continue

        if stripped.startswith('>'):
            quote_buf.append(stripped.lstrip('>').lstrip())
            continue
        flush_quote()

        if not stripped:
            continue

        m = re.match(r'^(\s*)-\s+\[([ xX])\]\s+(.+)$', line)
        if m:
            level = len(m.group(1)) // 2 + 1
            add_paragraph(doc, m.group(3), checkbox=(m.group(2).lower() == 'x'),
                          indent_level=level)
            continue

        m = re.match(r'^(\s*)[-*]\s+(.+)$', line)
        if m:
            level = len(m.group(1)) // 2 + 1
            add_paragraph(doc, m.group(2), bullet=True, indent_level=level)
            continue

        m = re.match(r'^(\s*)(\d+)\.\s+(.+)$', line)
        if m:
            level = len(m.group(1)) // 2 + 1
            add_paragraph(doc, m.group(3), number=m.group(2), indent_level=level)
            continue

        add_paragraph(doc, stripped)

    if in_code and code_buf:
        add_code_block(doc, '\n'.join(code_buf))
    flush_table()
    flush_quote()

    doc.save(docx_path)
    return docx_path


if __name__ == '__main__':
    import argparse
    parser = argparse.ArgumentParser(description='Markdown → Word docx 转换（含封面页 + 目录页）')
    parser.add_argument('md_path', nargs='?', help='输入 Markdown 文件路径')
    parser.add_argument('docx_path', nargs='?', help='输出 docx 文件路径（默认与 md 同名同目录）')
    parser.add_argument('--no-cover', action='store_true', help='不生成封面页')
    parser.add_argument('--no-toc', action='store_true', help='不生成目录页')
    parser.add_argument('--title', help='覆盖文档标题')
    parser.add_argument('--version', help='覆盖版本号')
    parser.add_argument('--organization', help='覆盖编制单位')
    parser.add_argument('--date', help='覆盖编制日期')
    args = parser.parse_args()

    if not args.md_path:
        # 兼容老用法：脚本同目录下找同名 md（按需修改）
        here = os.path.dirname(os.path.abspath(__file__))
        candidates = [f for f in os.listdir(here) if f.endswith('.md')]
        if not candidates:
            parser.print_help()
            raise SystemExit('未指定 md_path 且脚本目录下没有 .md 文件')
        args.md_path = os.path.join(here, candidates[0])

    if not args.docx_path:
        args.docx_path = os.path.splitext(args.md_path)[0] + '.docx'

    override = {k: v for k, v in {
        'title': args.title, 'version': args.version,
        'organization': args.organization, 'date': args.date,
    }.items() if v}

    out = convert(
        args.md_path, args.docx_path,
        with_cover=not args.no_cover,
        with_toc=not args.no_toc,
        metadata_override=override or None,
    )
    print('saved:', out)
